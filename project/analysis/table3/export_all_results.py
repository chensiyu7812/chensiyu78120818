"""Read-only aggregation of all measured v7 results into an Excel/CSV bundle."""
import csv
import hashlib
import json
import sqlite3
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from statistics import mean

import numpy as np

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
RUN = ROOT / 'outputs/table3_under10/run_identity_v7'
import argparse
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--output', type=Path, default=RUN / 'all_results_rebuilt')
OUT = parser.parse_args().output.resolve()
PLAN = ROOT / 'outputs/table3_rerun_v1/plan'
ARMS = ['pm_off', 'pm_on', 'rule_off', 'rule_on', 'fixed_off', 'fixed_on']
LABELS = ['PM OFF', 'PM ON', 'Rule OFF', 'Rule ON', 'Fixed OFF', 'Fixed ON']
QUALITY = ['overall', 'emotional_support', 'personalization', 'memory_appropriateness',
           'factual_grounding', 'temporal_consistency', 'non_intrusiveness']
RISK = ['selected_evidence_misuse', 'unnecessary_exposure', 'stale_or_conflict',
        'unsupported_personal_claim', 'source_set_appropriateness', 'strategy_overuse',
        'strategy_omission', 'response_support_sufficiency', 'overall_risk', 'omission_severity']

def read(path): return json.loads(Path(path).read_text())
def jsonl(path): return [json.loads(s) for s in Path(path).open()]
def digest(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()

source_paths = set()
def source(path):
    source_paths.add(Path(path))
    return Path(path)

manifest = read(source(RUN / 'result_manifest.json'))
freeze = read(source(ROOT / 'outputs/table3_under10/frozen_identity_v7/manifest.json'))
assert freeze['freeze_sha256'] == manifest['freeze_sha256']
for name, sha in manifest['hashes'].items():
    assert digest(source(RUN / name)) == sha
for name in ['arms.jsonl', 'units.jsonl']:
    path = source(PLAN / name)
    assert digest(path) == freeze['files'][str(path.relative_to(ROOT))]
plan = {(r['unit_id'], r['arm']): r for r in jsonl(PLAN / 'arms.jsonl')}
units = {r['unit_id']: r for r in jsonl(PLAN / 'units.jsonl')}
base = jsonl(RUN / 'observations.jsonl')
assert len(base) == len(plan) == 1224
assert {(r['unit_id'], r['arm']) for r in base} == set(plan)
db = sqlite3.connect(f'file:{source(RUN / "ledger.sqlite")}?mode=ro', uri=True)
db.row_factory = sqlite3.Row
attempts = [dict(r) for r in db.execute('SELECT * FROM attempts ORDER BY created_at,id')]
db.close()
assert not [a for a in attempts if a['status'] not in ('done', 'invalid')]
done = {}
for a in attempts:
    if a['status'] == 'done':
        assert a['job_id'] not in done
        done[a['job_id']] = a

def scores_for(phase):
    jobs = jsonl(source(RUN / f'{phase}.jobs.jsonl'))
    scores = {}
    for j in jobs:
        a = done[j['id']]
        p = json.loads(a['parsed_json'])
        if j['stage'] == 'quality':
            for candidate in p['candidates']:
                mapping = j['mapping'][candidate['candidate_id']]
                key = (j['unit_id'], mapping['condition'], 'quality', j['variant'])
                assert key not in scores
                scores[key] = dict(candidate, position=mapping['position'])
        else:
            key = (j['unit_id'], j['arm'], j['stage'], 0)
            assert key not in scores
            scores[key] = p
    return jobs, scores

main_jobs, main = scores_for('main')
assert len(main_jobs) == 2652 and len(main) == 3672
generations = {}
for a in attempts:
    if a['phase'] == 'generation' and a['status'] == 'done':
        j, p = json.loads(a['job_json']), json.loads(a['response_json'])
        key = (j['unit_id'], j['arm'])
        assert key not in generations
        generations[key] = (a['job_id'], p)
assert set(generations) == set(plan)
eos_summary = read(source(RUN / 'eos_generation_summary.json'))
eos = {}
for path in (ROOT / 'outputs/table3_under10/wire_amendment_v4/extended_replies').glob('*.json'):
    source(path)
    assert digest(path) == freeze['files'][str(path.relative_to(ROOT))]
    e = read(path)
    key = (e['unit_id'], e['arm'])
    assert key not in eos
    eos[key] = e
assert len(eos) == 28
rows = []
for original in base:
    r = dict(original)
    uid, arm = r['unit_id'], r['arm']
    key = uid, arm
    q, risk, omission = (main[(uid, arm, stage, 0)] for stage in ['quality', 'risk', 'omission'])
    assert all(r[k] == q[k] for k in QUALITY)
    assert r['misuse_severity'] == risk['selected_evidence_misuse']
    assert r['omission_severity'] == omission['omission_severity']
    r.update({f'risk_{k}': v for k, v in risk.items()})
    r.update({f'omission_{k}': v for k, v in omission.items()})
    r['quality_reason'] = q['reason']
    r['quality_candidate_position'] = q['position']
    p, u = plan[key], units[uid]
    r.update(seed=u['seed'], turn_index=u['turn_index'], requested_action_id=p['requested_action_id'],
             effective_action_id=p['effective_action_id'],
             candidate_items=len(p['candidate_memory']), kept_items=len(p['kept_memory']),
             dropped_items=len(p['dropped_memory_ids']), strategy_cards=len(p['selected_strategy']),
             requested_memory_sources=len(p['requested_memory_sources']),
             effective_memory_sources=len({m['source'] for m in p['kept_memory']}),
             offline_policy_compute_ms=p['offline_policy_compute_ms'],
             offline_filter_compute_ms=p['offline_filter_compute_ms'])
    for memory in ['MP', 'MS', 'ME']:
        r['requested_' + memory] = int(memory in p['requested_memory_sources'])
        r['candidate_' + memory + '_items'] = sum(m['source'] == memory for m in p['candidate_memory'])
        r['kept_' + memory + '_items'] = sum(m['source'] == memory for m in p['kept_memory'])
    r['requested_RS'] = int(bool(p['selected_strategy']))
    jid, g = generations[key]
    original_ids = g['local_backend']['completion_token_ids']
    r['generator_input_tokens'] = g['usage']['prompt_tokens']
    r['original_output_tokens'] = g['usage']['completion_tokens']
    r['original_generation_seconds'] = g['local_backend']['latency_seconds']
    r['original_finish_reason'] = g['choices'][0]['finish_reason']
    r['was_cap_hit'] = int(r['original_finish_reason'] == 'length')
    if key in eos:
        e = eos[key]
        assert e['original_job_id'] == jid and e['natural_stop'] and e['prefix_matches']
        assert e['completion_token_ids'][:100] == original_ids
        assert e['input_ids_sha256'] == g['local_backend']['input_ids_sha256']
        r['generator_output_tokens'] = e['completion_tokens']
        r['generation_seconds'] = e['latency_seconds']
        r['response_text'] = e['text']
        final_ids = e['completion_token_ids']
    else:
        assert r['original_finish_reason'] == 'stop'
        r['generator_output_tokens'] = g['usage']['completion_tokens']
        r['generation_seconds'] = g['local_backend']['latency_seconds']
        r['response_text'] = g['choices'][0]['message']['content']
        final_ids = original_ids
    assert final_ids[-1] in (128001, 128008, 128009)
    assert len(final_ids) == r['generator_output_tokens']
    r['text_changed_after_eos'] = int(r['response_text'] != g['choices'][0]['message']['content'])
    r['natural_stop'] = 1
    r['empty_memory'] = int(r['kept_items'] == 0)
    r['risk_auxiliary_guard_exception'] = int(r['risk_omission_severity'] != 0)
    r['generator_total_tokens'] = r['generator_input_tokens'] + r['generator_output_tokens']
    r['end_to_end_latency_ms'] = None
    r['retrieval_latency_ms'] = None
    rows.append(r)
assert sum(r['was_cap_hit'] for r in rows) == 28
assert sum(r['text_changed_after_eos'] for r in rows) == 24
assert sum(r['risk_auxiliary_guard_exception'] for r in rows) == 1
groups = {arm: [r for r in rows if r['arm'] == arm] for arm in ARMS}
assert all(len(rs) == 204 for rs in groups.values())

metrics = []
def add(group, label, unit, field=None, kind='mean', note='', fn=None):
    result = dict(category=group, metric=label, unit=unit, aggregation=kind, field=field or '')
    for arm, rs in groups.items():
        values = [r[field] for r in rs] if field else []
        if fn: value = fn(rs)
        elif kind == 'mean': value = mean(values)
        elif kind == 'sum': value = sum(values)
        elif kind == 'p50': value = float(np.quantile(values, .5))
        elif kind == 'p95': value = float(np.quantile(values, .95))
        elif kind == 'max': value = max(values)
        else: raise ValueError(kind)
        result[arm] = value
    result['definition'] = note
    metrics.append(result)

add('样本', '状态／回复数', '条', 'natural_stop', 'sum', '全部固定状态等权。')
for field, label in zip(QUALITY, ['Overall 总体质量 ↑', '情感支持 ↑', '个性化 ↑', '记忆适当性 ↑', '事实依据 ↑', '时间一致性 ↑', '非侵入性 ↑']):
    add('质量', label, '1–5', field, note='匿名六候选质量 judge；均值。')
for field, label, scale in [
    ('source_set_appropriateness', '来源集合适配度 ↑', '1–5'),
    ('selected_evidence_misuse', '上下文误用严重度 ↓', '0–3'),
    ('unnecessary_exposure', '不必要信息暴露 ↓', '0–3'),
    ('stale_or_conflict', '过时或冲突信息使用 ↓', '0–3'),
    ('unsupported_personal_claim', '无依据的个人信息断言 ↓', '0–3'),
    ('strategy_overuse', '策略过用 ↓', '0–3'),
    ('strategy_omission', '策略遗漏 ↓', '0–3'),
    ('response_support_sufficiency', '充分性：已选上下文 ↑', '1–5'),
    ('overall_risk', '总体风险 ↓', '0–3'),
]: add('风险审计', label, scale, 'risk_' + field, note='selected-context risk judge；均值。')
add('风险审计', '任意误用率（≥1）↓', '%', fn=lambda rs:100*mean(r['misuse_any'] for r in rs), note='分母为各组全部 204 条。')
add('风险审计', '明确误用率（≥2）↓', '%', fn=lambda rs:100*mean(r['misuse_clear'] for r in rs), note='分母为各组全部 204 条。')
for verdict, label in [('acceptable', 'Acceptable 数量'), ('minor_issue', 'Minor issue 数量'), ('major_issue', 'Major issue 数量'), ('uncertain', 'Uncertain 数量')]:
    add('风险审计', label, '条', kind='count', fn=lambda rs,v=verdict:sum(r['risk_verdict']==v for r in rs), note='直接统计 judge verdict，不能由各严重度阈值反推。')
add('遗漏审计', '遗漏严重度 ↓', '0–3', 'omission_omission_severity', note='独立 authorized-context omission judge；不用 risk 阶段同名辅助字段。')
add('遗漏审计', '充分性：完整授权上下文 ↑', '1–5', 'omission_response_support_sufficiency', note='独立 omission judge；与已选上下文充分性分开报告。')
add('遗漏审计', '任意遗漏率（≥1）↓', '%', fn=lambda rs:100*mean(r['omission_omission_severity']>=1 for r in rs))
add('遗漏审计', '明确遗漏率（≥2）↓', '%', fn=lambda rs:100*mean(r['omission_clear'] for r in rs))
add('遗漏审计', '明确遗漏数量（≥2）↓', '条', 'omission_clear', 'sum')
for field,label,unit in [
    ('source_invocations','请求来源总数（含 RS）↓','个/回复'),
    ('requested_memory_sources','请求 Memory 来源数 ↓','个/回复'),
    ('effective_memory_sources','保留 Memory 来源数','个/回复'),
    ('strategy_cards','Strategy 卡片数','张/回复'),
    ('candidate_items','候选 Memory 条目数','条/回复'),
    ('kept_items','保留 Memory 条目数','条/回复'),
    ('dropped_items','删除 Memory 条目数','条/回复'),
    ('candidate_memory_tokens_est','候选 Memory tokens ↓','估计 tokens'),
    ('memory_tokens_est','保留 Memory tokens ↓','估计 tokens'),
]: add('资源',label,unit,field,note='来源数为逻辑来源数；Memory token 沿用 v1 chars/4 估计。')
add('资源', 'Memory token 削减比例', '%', fn=lambda rs:100*(1-sum(r['memory_tokens_est'] for r in rs)/sum(r['candidate_memory_tokens_est'] for r in rs)),note='1 − 组内保留 tokens 总和 / 候选 tokens 总和；不是逐状态百分比的均值。')
add('资源', '空 Memory 回复数', '条', 'empty_memory', 'sum')
add('资源', '空 Memory 比例', '%', fn=lambda rs:100*mean(r['empty_memory'] for r in rs))
for source_name in ['MP', 'MS', 'ME', 'RS']:
    add('来源分解',source_name+' 请求比例','%',fn=lambda rs,s=source_name:100*mean(r['requested_'+s] for r in rs))
for memory in ['MP', 'MS', 'ME']:
    for prefix,verb in [('candidate','候选'),('kept','保留')]:
        add('来源分解',memory+' '+verb+'条目数','条/回复',f'{prefix}_{memory}_items')
for field, label in [('generator_input_tokens','完整生成输入'),('generator_output_tokens','最终生成输出')]:
    for kind,suffix in [('mean','均值'),('p50','P50'),('p95','P95'),('max','最大值')]:
        add('生成用量',label+' tokens '+suffix,'实际 tokens',field,kind,note='Llama tokenizer 实际计数；输出含终止 token，采用自然结束后的最终版本。')
add('生成用量','输入＋输出 tokens 均值','实际 tokens','generator_total_tokens')
add('完成检查','原 100-token 上限触及数','条','was_cap_hit','sum')
add('完成检查','补全后文本改变数','条','text_changed_after_eos','sum')
add('完成检查','最终自然结束数','条','natural_stop','sum')
add('完成检查','risk 辅助校验例外数','条','risk_auxiliary_guard_exception','sum',note='保留最早原始评分；不改分、不删样本。该字段不是遗漏结果指标。')
for kind,suffix in [('mean','均值'),('p50','P50'),('p95','P95'),('max','最大值')]:
    add('计时诊断','生成环节耗时 '+suffix,'秒','generation_seconds',kind,note='本地 model.generate 与解码附近计时；1196 条原批次＋28 条 EOS 重生成批次，非统一端到端 benchmark。')
add('计时诊断','原批次自然结束子集：生成均时','秒',kind='subset mean',fn=lambda rs:mean(r['generation_seconds'] for r in rs if not r['was_cap_hit']),note='不同组子集构成不同，不能据此筛选比较或替代全样本结果。')
add('计时诊断','EOS 补生成子集：生成均时','秒',kind='subset mean',fn=lambda rs:mean(r['generation_seconds'] for r in rs if r['was_cap_hit']),note='只对触及旧上限的样本测量；输出完整重生成耗时，不与原 100-token 耗时相加。')
for field,label in [('offline_policy_compute_ms','离线 policy.choose 耗时'),('offline_filter_compute_ms','离线 Filter 函数耗时')]:
    for kind,suffix in [('mean','均值'),('p50','P50'),('p95','P95')]:
        add('计时诊断',label+' '+suffix,'毫秒',field,kind,note='来自离线规划，含首次调用等现场开销；不含完整线上检索链路。OFF Filter 也调用同一函数用于诊断。')
add('未测量','完整检索耗时','毫秒',kind='unavailable',fn=lambda rs:None,note='未保存统一计时，不能从其他耗时相减推算。')
add('未测量','端到端回复延迟','毫秒',kind='unavailable',fn=lambda rs:None,note='未统一测量 PM＋检索＋Filter＋生成；不能将离线与生成计时直接相加作为线上延迟。')

pilot_rows, pilot_summary = [], []
for phase in ['pilot_mini','pilot_reference']:
    jobs, observations = scores_for(phase)
    assert len(jobs)==168 and len(observations)==288
    for (uid,arm,stage,variant), score in observations.items():
        pilot_rows.append(dict(phase=phase,unit_id=uid,arm=arm,stage=stage,variant=variant,**score))
    for arm in ARMS:
        for stage,variants in [('quality',[0,1]),('risk',[0]),('omission',[0])]:
            for variant in variants:
                rs=[p for p in pilot_rows if p['phase']==phase and p['arm']==arm and p['stage']==stage and p['variant']==variant]
                assert len(rs)==12
                for field in (QUALITY if stage=='quality' else RISK if stage=='risk' else ['omission_severity','response_support_sufficiency']):
                    pilot_summary.append(dict(phase=phase,arm=arm,stage=stage,variant=variant,metric=field,n=12,mean=mean(r[field] for r in rs)))
gate = read(source(RUN/'pilot_gate.json'))
gate_rows = []
def flatten(value,prefix=''):
    if isinstance(value,dict):
        for key,v in value.items():flatten(v,prefix+'.'+key if prefix else key)
    else:gate_rows.append({'item':prefix,'value':value})
flatten(gate)

budget = []
for phase in sorted({a['phase'] for a in attempts}):
    for status in sorted({a['status'] for a in attempts if a['phase']==phase}):
        rs=[a for a in attempts if a['phase']==phase and a['status']==status]
        budget.append(dict(phase=phase,status=status,attempts=len(rs),ledger_usd_upper=sum(a['actual_nano'] or 0 for a in rs)/1e9,note='包含历史评分与重试；不是各方法的部署成本。'))
ledger_total=sum(a['actual_nano'] or 0 for a in attempts)/1e9
assert abs(ledger_total-manifest['budget']['charged_upper_usd'])<1e-10
budget.append(dict(phase='TOTAL_LEDGER',status='settled',attempts=len(attempts),ledger_usd_upper=ledger_total,note='账本未缓存单价口径。'))
diagnostic=read(source(ROOT/'outputs/table3_under10/identity_amendment_v6/incident.json'))['diagnostic_probe']['cost_usd_approx']
budget.append(dict(phase='OFF_LEDGER_DIAGNOSTIC',status='estimate',attempts=1,ledger_usd_upper=diagnostic,note='事件记录估计；尚未按原始 usage 对账，不计入任何评分。'))
budget.append(dict(phase='TOTAL_WITH_DIAGNOSTIC',status='estimate',attempts=len(attempts)+1,ledger_usd_upper=ledger_total+diagnostic,note='含诊断估计，非精确官方账单。'))
attempt_export=[]
active_ids={j['id'] for j in main_jobs}
for a in attempts:
    j=json.loads(a['job_json']); resp=json.loads(a['response_json']) if a['response_json'] else {}
    usage=resp.get('usage',{})
    attempt_export.append(dict(attempt_id=a['id'],phase=a['phase'],stage=j['stage'],arm=j.get('arm'),
        unit_id=j['unit_id'],attempt_number=a['number'],status=a['status'],active_main_score=int(a['job_id'] in active_ids and a['status']=='done'),
        actual_usd=(a['actual_nano']/1e9 if a['actual_nano'] is not None else None),prompt_tokens=usage.get('prompt_tokens'),
        completion_tokens=usage.get('completion_tokens'),cached_tokens=usage.get('prompt_tokens_details',{}).get('cached_tokens'),created_at=a['created_at']))
intervals=read(RUN/'paired_intervals.json')
actions=[]
for arm,rs in groups.items():
    for field in ['requested_action_id','effective_action_id']:
        for action,count in sorted(Counter(r[field] for r in rs).items()):
            actions.append(dict(arm=arm,action_type=field,action=action,count=count,percent=100*count/204))
histograms=[]
for arm,rs in groups.items():
    fields=QUALITY+['risk_'+k for k in RISK]+['omission_omission_severity','omission_response_support_sufficiency']
    for field in fields:
        for value,count in sorted(Counter(r[field] for r in rs).items()):
            histograms.append(dict(arm=arm,metric=field,score=value,count=count,percent=100*count/204))

OUT.mkdir(parents=True, exist_ok=True)
def write_csv(name, records):
    fields=list(dict.fromkeys(k for r in records for k in r))
    with (OUT/name).open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=fields);writer.writeheader();writer.writerows(records)
    return fields

datasets=[('总表','all_metrics.csv',metrics),('逐条数据','per_response.csv',rows),
          ('配对区间','paired_intervals.csv',intervals),('评分分布','score_histograms.csv',histograms),
          ('来源动作分布','action_distribution.csv',actions),('Pilot均值','pilot_summary.csv',pilot_summary),
          ('Pilot逐条','pilot_observations.csv',pilot_rows),('Pilot门槛','pilot_gate.csv',gate_rows),
          ('费用汇总','budget.csv',budget),('全部尝试账目','attempts.csv',attempt_export)]
wb=Workbook();ws=wb.active;ws.title='说明'
notes=[
    ('范围','本次 v7 六组完整重跑；204 状态、34 场景、18 用户。主表为 GPT-4.1-mini，GPT-4o 只在 Pilot 页。'),
    ('总表读法','指标为行、六组为列；除数量、百分比和 P50/P95/max 外均为状态等权均值。完整精度在 CSV/Excel。'),
    ('质量量表','质量、来源适配度与两种充分性为 1–5；风险与遗漏严重度为 0–3。百分比列采用 0–100 数值。'),
    ('两个充分性','risk 的 response_support_sufficiency 针对已选上下文；omission 的同名字段针对完整授权上下文，分列保留。'),
    ('风险辅助字段','risk.omission_severity 是应为 0 的校验字段，只有一条例外，逐条保留但不作为遗漏指标。'),
    ('来源与 token','逻辑来源包含 RS；Memory tokens 为 chars/4 估计，生成输入输出为 Llama tokenizer 实数，两者不互相替换。'),
    ('计时范围','生成耗时包括本地 generate 与解码附近计时；没有 PM、检索、Filter 或完整服务链路。1196 条原批次＋28 条 EOS 补生成，全部保留。'),
    ('离线计时','来自规划阶段，policy.choose 与 Filter 函数各计时；首次调用等现场开销不剔除。不能与生成耗时简单拼成端到端 benchmark。'),
    ('生成完整性','原 28 条截断全部补至 EOS，前 100 token 匹配，24 条文本改变。输入 token 未改变。'),
    ('配对区间','保留既有 238 条配对区间：7 对比×17 指标×2 聚类；用户为主、场景为敏感性。未为本次新增汇总字段追加检验。'),
    ('单条例外','fixed_off 中 1 条 risk 辅助校验例外，采用最早原始响应，未改分、未删样本；原文五次 reason 相同的说法已有更正。'),
    ('费用','账本包含历史评分及失败尝试，另有约 $0.0013 诊断估计。费用不是部署延迟，也不按 arm 均摊匿名六候选评分。'),
    ('论文呈现','本文件先完整整理已测量结果，不选取有利指标或据此调整协议。论文选列与补测延迟留待下一步讨论。'),
    ('来源冻结',manifest['freeze_sha256']),
]
for r in notes:ws.append(r)
for title,name,records in datasets:
    if title=='总表':
        order=['category','metric','unit',*ARMS,'definition','aggregation','field']
        records=[{k:r[k] for k in order} for r in records]
    fields=write_csv(name,records)
    ws=wb.create_sheet(title)
    headers={'category':'类别','metric':'指标','unit':'单位','definition':'定义与测量范围',
             'aggregation':'汇总方式','field':'原始字段',**dict(zip(ARMS,LABELS))}
    ws.append([headers.get(k,k) for k in fields] if title=='总表' else fields)
    for record in records:
        ws.append([record.get(k) for k in fields])
    ws.freeze_panes='D2' if title=='总表' else 'C2'
    ws.auto_filter.ref=ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value,float):cell.number_format='0.0000'
            elif isinstance(cell.value,str) and cell.value.startswith(('=','+','-','@')):cell.data_type='s'
    if title=='总表':
        previous=None
        for row in ws.iter_rows(min_row=2):
            category=row[0].value
            if category!=previous:
                for cell in row:cell.fill=PatternFill('solid',fgColor='E5EDF5')
            previous=category
for ws in wb:
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='24435B')
    for col in ws.columns:
        letter=get_column_letter(col[0].column)
        max_len=max(len(str(c.value or '')) for c in list(col)[:100])
        ws.column_dimensions[letter].width=min(55,max(14,max_len+2))
    ws.sheet_view.showGridLines=False
wb['总表'].column_dimensions['B'].width=43
wb['总表'].column_dimensions['A'].width=16
wb['总表'].column_dimensions['C'].width=15
for letter in 'DEFGHI':wb['总表'].column_dimensions[letter].width=16
wb['总表'].column_dimensions['J'].width=82
wb['说明'].column_dimensions['A'].width=22
wb['说明'].column_dimensions['B'].width=118
for row in wb['说明']:
    row[1].alignment=Alignment(wrap_text=True,vertical='top')
    wb['说明'].row_dimensions[row[0].row].height=40
workbook=OUT/'Table3_All_Results.xlsx';wb.save(workbook)

def fmt(value,unit):
    if value is None:return '未测'
    if unit=='条' or (unit=='实际 tokens' and float(value).is_integer()):return str(int(value))
    if unit=='%':return f'{value:.2f}%'
    if unit in ['实际 tokens','估计 tokens']:return f'{value:.1f}'
    return f'{value:.3f}'
table_lines=['| 类别／指标 | 单位 | '+' | '.join(LABELS)+' |','|---|---|'+'---:|'*6]
for m in metrics:
    table_lines.append('| '+m['category']+'／'+m['metric']+' | '+m['unit']+' | '+' | '.join(fmt(m[a],m['unit']) for a in ARMS)+' |')
md=['# Table III 全部结果总表','',
    '本次六组各 204 条，指标为行、方法为列。所有正式数值来自同一 GPT-4.1-mini 主评分；GPT-4o pilot、配对区间、费用、原始分数和解释分别附在工作簿中。本表不选择论文主指标。','',
    *table_lines,'',
    '计时是现场生成及离线规划诊断，范围见工作簿说明；端到端延迟未测。两个充分性分别来自已选上下文 risk 与完整授权上下文 omission，不能混为同一评分。数量和百分比以各组 204 为分母。Memory token 为估计，生成 token 为实际计数。','',
    f'工作簿含 {len(metrics)} 行汇总（其中 2 行标明未测量）、1,224 行逐条记录、{len(intervals)} 条既有配对区间、{len(pilot_rows)} 条 pilot 观测及全部 {len(attempt_export)} 次账本尝试。', '',
    '[Excel 工作簿](Table3_All_Results.xlsx) · [总表 CSV](all_metrics.csv) · [逐条数据 CSV](per_response.csv) · [例外说明](../FINAL_AUDIT_CN.md)','']
(OUT/'ALL_RESULTS_CN.md').write_text('\n'.join(md))
(OUT/'metric_definitions.json').write_text(json.dumps(metrics,ensure_ascii=False,indent=2)+'\n')
qa={'status':'PASS','created_at':datetime.now(timezone.utc).isoformat(),'freeze_sha256':manifest['freeze_sha256'],
    'main_rows':len(rows),'main_calls':len(main_jobs),'summary_rows':len(metrics),'paired_intervals':len(intervals),
    'pilot_observations':len(pilot_rows),'ledger_attempts':len(attempts),'ledger_usd_upper':ledger_total,
    'eos_replacements':len(eos),'source_sha256':{str(p):digest(p) for p in sorted(source_paths)},
    'script_sha256':digest(__file__),'checks':['Frozen source/result hashes','Unique active score association',
    'Quality and risk/omission values reconcile to original export','Six complete paired groups',
    'EOS prefix/input/final-stop checks','Ledger total including historical attempts','Workbook reopened with expected rows']}
check=load_workbook(workbook,read_only=True,data_only=True)
assert check['总表'].max_row==len(metrics)+1 and check['逐条数据'].max_row==1225
assert check['Pilot逐条'].max_row==577
check.close()
qa['output_sha256']={p.name:digest(p) for p in OUT.iterdir() if p.is_file() and p.name!='QA_MANIFEST.json'}
(OUT/'QA_MANIFEST.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({k:qa[k] for k in ['status','main_rows','summary_rows','paired_intervals','pilot_observations','ledger_attempts']},ensure_ascii=False))
print('\n'.join(table_lines))
